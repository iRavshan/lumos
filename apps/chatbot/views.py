import json
import time
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import ChatbotConfig, ChatSession, ChatMessage
from django.db.models import Q, Count
from django.urls import reverse
from .forms import ChatbotConfigForm, TelegramBotConfigForm, AgentConfigForm
from .services import generate_rag_response, extract_contact_info, analyze_session_insights
from .telegram_service import (
    verify_telegram_bot_token,
    set_telegram_webhook,
    delete_telegram_webhook,
    handle_telegram_update,
)


def cors_json_response(data, status=200):
    response = JsonResponse(data, status=status)
    response['Access-Control-Allow-Origin'] = '*'
    response['Access-Control-Allow-Methods'] = 'GET, POST, OPTIONS'
    response['Access-Control-Allow-Headers'] = 'Content-Type, X-Requested-With'
    return response


@csrf_exempt
def api_chatbot_config(request, api_key):
    if request.method == 'OPTIONS':
        return cors_json_response({})

    chatbot = get_object_or_404(ChatbotConfig, api_key=api_key)
    # Build favicon proxy URL for the JS widget
    from django.urls import reverse
    logo_url = reverse('businesses:favicon_proxy', kwargs={'business_id': chatbot.business.id})
    return cors_json_response({
        'status': 'success',
        'is_active': chatbot.is_active,
        'bot_name': chatbot.bot_name,
        'business_name': chatbot.business.name,
        'welcome_message': chatbot.welcome_message,
        'theme_color': chatbot.theme_color,
        'suggested_questions': [],
        'business_logo_url': logo_url if (chatbot.business.website or chatbot.business.telegram) else None,
        'response_delay_enabled': chatbot.response_delay_enabled,
        'response_delay_seconds': chatbot.first_message_delay_seconds,
        'first_message_delay_seconds': chatbot.first_message_delay_seconds,
        'subsequent_message_delay_seconds': chatbot.subsequent_message_delay_seconds,
        'split_messages': chatbot.split_messages,
    })


@csrf_exempt
def api_chat_message(request, api_key):
    if request.method == 'OPTIONS':
        return cors_json_response({})

    if request.method != 'POST':
        return cors_json_response({'error': 'Faqat POST so\'rov qabul qilinadi.'}, status=405)

    try:
        chatbot = ChatbotConfig.objects.select_related('business').get(api_key=api_key)
    except ChatbotConfig.DoesNotExist:
        return cors_json_response({'error': 'Bunday chatbot topilmadi.'}, status=404)

    if not chatbot.is_active:
        return cors_json_response({'error': 'Chatbot hozirda nofaol holatda.'}, status=403)

    try:
        body = json.loads(request.body.decode('utf-8'))
    except Exception:
        body = request.POST

    user_message = body.get('message', '').strip()
    session_id = body.get('session_id', '').strip()

    if not user_message:
        return cors_json_response({'error': 'Xabar matni bo\'sh bo\'lishi mumkin emas.'}, status=400)

    # Get client IP
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        ip = x_forwarded_for.split(',')[0].strip()
    else:
        ip = request.META.get('REMOTE_ADDR')

    # Get or create chat session
    session = None
    if session_id:
        session, created = ChatSession.objects.get_or_create(
            chatbot=chatbot,
            session_id=session_id,
            defaults={'visitor_ip': ip}
        )
        if not session.visitor_ip and ip:
            session.visitor_ip = ip

        # Automatically extract and update contact details if provided by visitor
        contacts = extract_contact_info(user_message)
        changed = False
        if contacts.get('phone') and not session.visitor_phone:
            session.visitor_phone = contacts['phone']
            changed = True
        if contacts.get('email') and not session.visitor_email:
            session.visitor_email = contacts['email']
            changed = True
        if contacts.get('name') and not session.visitor_name:
            session.visitor_name = contacts['name']
            changed = True

        if changed or not created:
            session.save()

    # Save user message
    user_msg = ChatMessage.objects.create(
        chatbot=chatbot,
        session=session,
        role='user',
        content=user_message
    )
    from .realtime import notify_chat_update
    notify_chat_update(chatbot.business_id, session.session_id if session else None, user_msg.id)

    try:
        # Generate RAG reply
        reply = generate_rag_response(chatbot, user_message)

        # Check if this chat requires human operator escalation
        from apps.team.services import check_if_needs_escalation, assign_session_to_operator
        needs_esc, reason = check_if_needs_escalation(user_message, reply)

        staff_assigned = None
        if session and (needs_esc or session.is_escalated):
            if not session.is_escalated:
                staff_assigned = assign_session_to_operator(session, reason)
                reply += f"\n\n👨‍💼 Savolingiz navbatchi mutaxassisimiz ({staff_assigned.name if staff_assigned else 'Operator'}) ga yo'naltirildi. Tez orada javob beriladi."

    except Exception as e:
        import logging
        logging.getLogger(__name__).error("Chat API xatolik: %s", e, exc_info=True)
        reply = "Kechirasiz, texnik nosozlik yuz berdi. Iltimos, qayta urinib ko'ring yoki biz bilan bevosita bog'laning."

    # Save assistant message
    asst_msg = ChatMessage.objects.create(
        chatbot=chatbot,
        session=session,
        role='assistant',
        content=reply
    )
    notify_chat_update(chatbot.business_id, session.session_id if session else None, asst_msg.id)

    # Split messages if enabled
    message_parts = [reply]
    if chatbot.split_messages and ('\n\n' in reply or '. ' in reply):
        raw_parts = [p.strip() for p in reply.split('\n\n') if p.strip()]
        if len(raw_parts) > 1:
            message_parts = raw_parts

    user_msg_count = session.messages.filter(role='user').count() if session else 1
    if chatbot.response_delay_enabled:
        active_delay = chatbot.first_message_delay_seconds if user_msg_count <= 1 else chatbot.subsequent_message_delay_seconds
    else:
        active_delay = 0

    return cors_json_response({
        'status': 'success',
        'reply': reply,
        'parts': message_parts,
        'delay_seconds': active_delay,
        'first_message_delay_seconds': chatbot.first_message_delay_seconds,
        'subsequent_message_delay_seconds': chatbot.subsequent_message_delay_seconds,
        'split_messages': chatbot.split_messages,
        'bot_name': chatbot.bot_name,
        'business_name': chatbot.business.name,
        'is_escalated': bool(session and session.is_escalated),
    })


@csrf_exempt
def api_poll_widget(request, api_key):
    """
    Allows widget to poll for any new messages (e.g. sent by Operator/Supervisor from Telegram Mini App).
    """
    if request.method == 'OPTIONS':
        return cors_json_response({})

    session_id = request.GET.get('session_id', '').strip()
    after_id = request.GET.get('after_id', '0')

    try:
        after_id = int(after_id)
    except ValueError:
        after_id = 0

    if not session_id:
        return cors_json_response({'messages': []})

    messages_qs = ChatMessage.objects.filter(
        chatbot__api_key=api_key,
        session__session_id=session_id,
        id__gt=after_id
    ).order_by('created_at')

    data = []
    for m in messages_qs:
        sender_name = m.sender_staff.name if m.sender_staff else (m.chatbot.bot_name if m.role == 'assistant' else 'Siz')
        data.append({
            'id': m.id,
            'role': m.role,
            'content': m.content,
            'sender': sender_name,
            'created_at': m.created_at.strftime("%H:%M")
        })

    return cors_json_response({
        'status': 'success',
        'messages': data
    })


from django.views.decorators.clickjacking import xframe_options_exempt


@xframe_options_exempt
def widget_iframe_view(request, api_key):
    chatbot = get_object_or_404(ChatbotConfig, api_key=api_key)
    return render(request, 'chatbot/widget_iframe.html', {
        'chatbot': chatbot,
    })


def demo_page_view(request, api_key):
    chatbot = get_object_or_404(ChatbotConfig, api_key=api_key)
    return render(request, 'chatbot/demo_page.html', {
        'chatbot': chatbot
    })


@login_required
def chatbot_settings_view(request):
    if not hasattr(request.user, 'business'):
        messages.warning(request, "Avval biznesingizni ro'yxatdan o'tkazing.")
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot, created = ChatbotConfig.objects.get_or_create(
        business=business,
        defaults={
            'bot_name': f"{business.name} AI",
            'welcome_message': f"Assalomu alaykum! «{business.name}» virtual yordamchisiman. Sizga qanday yordam bera olaman?"
        }
    )

    if request.method == 'POST':
        form = ChatbotConfigForm(request.POST, instance=chatbot)
        if form.is_valid():
            form.save()
            messages.success(request, "Chatbot sozlamalari muvaffaqiyatli saqlandi!")
            return redirect('businesses:dashboard')
    else:
        form = ChatbotConfigForm(instance=chatbot)

    return render(request, 'chatbot/settings.html', {
        'form': form,
        'chatbot': chatbot,
        'business': business
    })


@login_required
def agent_settings_view(request):
    if not hasattr(request.user, 'business'):
        messages.warning(request, "Avval biznesingizni ro'yxatdan o'tkazing.")
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot, created = ChatbotConfig.objects.get_or_create(
        business=business,
        defaults={
            'bot_name': f"{business.name} AI",
            'welcome_message': f"Assalomu alaykum! «{business.name}» virtual yordamchisiman. Sizga qanday yordam bera olaman?"
        }
    )

    if request.method == 'POST':
        form = AgentConfigForm(request.POST, instance=chatbot)
        website = request.POST.get('website', '').strip()
        description = request.POST.get('description', '').strip()

        if form.is_valid():
            form.save()

            # Business website and description update if provided
            updated_business = False
            if website and website != (business.website or ''):
                business.website = website
                updated_business = True
            if description and description != (business.description or ''):
                business.description = description
                updated_business = True
            if updated_business:
                business.save()

            messages.success(request, "Savdo agenti sozlamalari muvaffaqiyatli saqlandi!")
            return redirect('chatbot:agent_settings')
    else:
        form = AgentConfigForm(instance=chatbot)

    return render(request, 'chatbot/agent_settings.html', {
        'form': form,
        'chatbot': chatbot,
        'business': business
    })


@login_required
def inbox_view(request, session_id=None):
    if not hasattr(request.user, 'business'):
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot = getattr(business, 'chatbot_config', None)

    if not chatbot:
        chatbot, _ = ChatbotConfig.objects.get_or_create(
            business=business,
            defaults={'bot_name': f"{business.name} AI"}
        )

    # Base sessions queryset with prefetched messages and their staff senders
    sessions_qs = chatbot.sessions.prefetch_related('messages__sender_staff').all()

    # Search filter
    q = request.GET.get('q', '').strip()
    if q:
        sessions_qs = sessions_qs.filter(
            Q(visitor_name__icontains=q) |
            Q(visitor_phone__icontains=q) |
            Q(visitor_email__icontains=q) |
            Q(session_id__icontains=q) |
            Q(messages__content__icontains=q)
        ).distinct()

    # Channel filter (all, website, telegram)
    channel_filter = request.GET.get('channel', 'all')
    if channel_filter == 'telegram':
        sessions_qs = sessions_qs.filter(session_id__startswith='tg_')
    elif channel_filter == 'website':
        sessions_qs = sessions_qs.exclude(session_id__startswith='tg_')

    # Counts
    all_sessions = chatbot.sessions.all()
    counts = {
        'all': all_sessions.count(),
        'website': all_sessions.exclude(session_id__startswith='tg_').count(),
        'telegram': all_sessions.filter(session_id__startswith='tg_').count(),
    }

    sessions_list = list(sessions_qs)

    # Active session selection & Insights (do not auto-open first chat on load)
    active_session = None
    insights = None
    if session_id:
        active_session = chatbot.sessions.filter(session_id=session_id).prefetch_related('messages__sender_staff', 'assigned_staff').first()
        if active_session:
            insights = analyze_session_insights(active_session)

    return render(request, 'chatbot/inbox.html', {
        'business': business,
        'chatbot': chatbot,
        'sessions': sessions_list,
        'active_session': active_session,
        'insights': insights,
        'counts': counts,
        'channel_filter': channel_filter,
        'search_query': q,
    })


@login_required
def update_lead_view(request, session_id):
    if not hasattr(request.user, 'business'):
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot = getattr(business, 'chatbot_config', None)
    session = get_object_or_404(ChatSession, chatbot=chatbot, session_id=session_id)

    if request.method == 'POST':
        session.visitor_name = request.POST.get('visitor_name', '').strip()
        session.visitor_phone = request.POST.get('visitor_phone', '').strip()
        session.visitor_email = request.POST.get('visitor_email', '').strip()
        session.status = request.POST.get('status', session.status)
        session.notes = request.POST.get('notes', '').strip()
        session.save()
        messages.success(request, f"«{session.display_name}» ma'lumotlari muvaffaqiyatli yangilandi.")

    return redirect('chatbot:inbox_detail', session_id=session_id)


@login_required
def send_inbox_message_view(request, session_id):
    if not hasattr(request.user, 'business'):
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot = getattr(business, 'chatbot_config', None)
    session = get_object_or_404(ChatSession, chatbot=chatbot, session_id=session_id)

    is_ajax = (
        request.headers.get('X-Requested-With') == 'XMLHttpRequest' or
        request.headers.get('Accept') == 'application/json' or
        ('application/json' in request.content_type if hasattr(request, 'content_type') and request.content_type else False)
    )

    if request.method == 'POST':
        raw_content = request.POST.get('content', '').strip()
        if not raw_content and request.body:
            try:
                data = json.loads(request.body.decode('utf-8'))
                raw_content = data.get('content', '').strip()
            except Exception:
                pass

        if raw_content:
            from .telegram_service import sanitize_for_telegram, send_telegram_bot_message
            from .realtime import notify_chat_update
            clean_content = sanitize_for_telegram(raw_content)

            # Create staff/admin reply
            msg = ChatMessage.objects.create(
                chatbot=chatbot,
                session=session,
                role='staff',
                content=clean_content
            )
            session.save() # updates last_message_at

            notify_chat_update(business.id, session.session_id, msg.id)

            # If telegram session, forward reply directly to Telegram chat
            if session.session_id.startswith('tg_') and chatbot.telegram_bot_token:
                try:
                    tg_chat_id = session.session_id.replace('tg_', '')
                    send_telegram_bot_message(
                        chatbot.telegram_bot_token,
                        tg_chat_id,
                        clean_content,
                        parse_mode='HTML'
                    )
                except Exception:
                    pass

            if is_ajax:
                return JsonResponse({
                    'status': 'success',
                    'message': {
                        'id': msg.id,
                        'role': msg.role,
                        'content': msg.formatted_content,
                        'sender_name': f"{business.name} (Operator)",
                        'created_at': msg.created_at.strftime("%H:%M"),
                        'date_divider': msg.date_divider,
                    },
                    'session': {
                        'id': session.session_id,
                        'display_name': session.display_name,
                        'last_message_at': session.last_message_at.strftime("%d %b"),
                        'messages_count': session.messages.count(),
                    }
                })

    if is_ajax:
        return JsonResponse({'status': 'error', 'message': 'Xabar matni bo\'sh bo\'lishi mumkin emas'}, status=400)

    return redirect('chatbot:inbox_detail', session_id=session_id)


@login_required
def inbox_sync_api(request):
    """
    Real-time synchronization endpoint for inbox.html.
    Guaranteed real-time delivery for Telegram, Web Widget, and TMA.
    """
    if not hasattr(request.user, 'business'):
        return JsonResponse({'error': 'Unauthorized'}, status=401)

    business = request.user.business
    chatbot = getattr(business, 'chatbot_config', None)
    if not chatbot:
        return JsonResponse({'status': 'ok', 'messages': [], 'sessions': []})

    from django.core.cache import cache

    active_session_id = request.GET.get('active_session_id', '').strip()
    try:
        after_id = int(request.GET.get('after_id', 0) or 0)
    except (ValueError, TypeError):
        after_id = 0

    try:
        client_v = int(request.GET.get('v', 0) or 0)
    except (ValueError, TypeError):
        client_v = 0

    # 1. Fetch new messages for active session directly by integer ID
    new_messages_data = []
    if active_session_id:
        active_session = chatbot.sessions.filter(session_id=active_session_id).first()
        if active_session:
            messages_qs = active_session.messages.filter(id__gt=after_id).select_related('sender_staff').order_by('created_at')
            for m in messages_qs:
                if m.role == 'user':
                    sender = 'Mijoz'
                elif m.role == 'staff':
                    sender = f"{m.sender_staff.name if m.sender_staff else business.name} (Operator)"
                else:
                    sender = chatbot.bot_name

                new_messages_data.append({
                    'id': m.id,
                    'role': m.role,
                    'content': m.formatted_content,
                    'sender_name': sender,
                    'created_at': m.created_at.strftime("%H:%M"),
                    'date_divider': m.date_divider,
                })

    # 2. Check business version from Redis for sidebar updates
    biz_key = f"biz:v:{business.id}"
    try:
        current_v = cache.get(biz_key)
    except Exception:
        current_v = None

    if current_v is None:
        current_v = 1
        try:
            cache.set(biz_key, current_v, timeout=86400)
        except Exception:
            pass
    else:
        try:
            current_v = int(current_v)
        except (ValueError, TypeError):
            current_v = 1

    sessions_data = []
    # Only serialize sidebar sessions if version changed or initial load (client_v == 0)
    if client_v == 0 or client_v < current_v or len(new_messages_data) > 0:
        recent_sessions = chatbot.sessions.prefetch_related('messages').order_by('-last_message_at')[:40]
        for s in recent_sessions:
            last_msg = s.last_message
            last_snippet = ''
            if last_msg and last_msg.content:
                clean_snip = last_msg.content.replace('\n', ' ').strip()
                last_snippet = (clean_snip[:45] + '...') if len(clean_snip) > 45 else clean_snip

            sessions_data.append({
                'session_id': s.session_id,
                'display_name': s.display_name,
                'visitor_phone': s.visitor_phone or '',
                'last_message_at': s.last_message_at.strftime("%d %b") if s.last_message_at else '',
                'messages_count': s.messages.count(),
                'last_snippet': last_snippet,
                'is_telegram': s.session_id.startswith('tg_'),
            })

    return JsonResponse({
        'status': 'updated' if (new_messages_data or sessions_data) else 'no_change',
        'v': current_v,
        'messages': new_messages_data,
        'sessions': sessions_data,
    })


@login_required
def chatbot_history_view(request):
    # Redirect legacy history link directly to the new powerful inbox view
    return redirect('chatbot:inbox')


@login_required
def telegram_bot_settings_view(request):
    if not hasattr(request.user, 'business'):
        messages.warning(request, "Avval biznesingizni ro'yxatdan o'tkazing.")
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot, _ = ChatbotConfig.objects.get_or_create(
        business=business,
        defaults={'bot_name': f"{business.name} AI"}
    )

    webhook_url = request.build_absolute_uri(
        reverse('chatbot:telegram_webhook', kwargs={'api_key': chatbot.api_key})
    )

    if request.method == 'POST':
        form = TelegramBotConfigForm(request.POST, instance=chatbot)
        if form.is_valid():
            token = form.cleaned_data.get('telegram_bot_token', '').strip()
            is_active = form.cleaned_data.get('telegram_bot_active', False)

            if token:
                # Verify token with Telegram
                is_valid, bot_info, err = verify_telegram_bot_token(token)
                if is_valid:
                    chatbot.telegram_bot_token = token
                    chatbot.telegram_bot_username = bot_info.get('username')
                    chatbot.telegram_bot_name = bot_info.get('first_name')
                    chatbot.telegram_bot_active = is_active

                    # Set webhook on Telegram (only for valid public HTTPS domains)
                    if is_active:
                        if webhook_url.startswith('https://') and 'localhost' not in webhook_url and '127.0.0.1' not in webhook_url:
                            wh_res = set_telegram_webhook(token, webhook_url)
                            if not wh_res.get('ok'):
                                delete_telegram_webhook(token)
                        else:
                            delete_telegram_webhook(token)
                    else:
                        delete_telegram_webhook(token)

                    chatbot.save()
                    messages.success(request, f"«@{chatbot.telegram_bot_username}» Telegram boti muvaffaqiyatli ulandi va sozlandi!")
                    return redirect('chatbot:telegram_settings')
                else:
                    form.add_error('telegram_bot_token', f"Telegram xatosi: {err}")
            else:
                # If user cleared token, delete webhook and deactivate
                if chatbot.telegram_bot_token:
                    delete_telegram_webhook(chatbot.telegram_bot_token)
                chatbot.telegram_bot_token = None
                chatbot.telegram_bot_username = None
                chatbot.telegram_bot_name = None
                chatbot.telegram_bot_active = False
                chatbot.save()
                messages.info(request, "Telegram bot ulanishi uzildi.")
                return redirect('chatbot:telegram_settings')
    else:
        form = TelegramBotConfigForm(instance=chatbot)

    return render(request, 'chatbot/telegram_settings.html', {
        'form': form,
        'chatbot': chatbot,
        'business': business,
        'webhook_url': webhook_url
    })


@csrf_exempt
def api_telegram_webhook(request, api_key):
    if request.method != 'POST':
        return HttpResponse("Only POST allowed", status=405)

    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception:
        return JsonResponse({'ok': False, 'error': 'Invalid JSON'}, status=400)

    result = handle_telegram_update(api_key, data)
    return JsonResponse(result)


@login_required
def export_inbox_excel_view(request):
    """
    Exports all leads/chat sessions into an Excel (.xlsx) file with two separate sheets:
    1. Vebsayt Murojaatlari (Website inquiries)
    2. Telegram Murojaatlari (Telegram bot inquiries)
    """
    if not hasattr(request.user, 'business'):
        return redirect('businesses:onboarding')

    business = request.user.business
    chatbot = getattr(business, 'chatbot_config', None)
    if not chatbot:
        messages.warning(request, "Chatbot ma'lumotlari topilmadi.")
        return redirect('chatbot:inbox')

    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.utils import timezone

    wb = openpyxl.Workbook()

    # Define styles
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid") # Indigo 600
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "№", 
        "Murojaatchi Ismi", 
        "Telefon Raqami", 
        "Email", 
        "Birinchi Murojaat", 
        "So'nggi Xabar Vaqti", 
        "Xabarlar Soni", 
        "Eslatma / Izoh"
    ]

    all_sessions = chatbot.sessions.prefetch_related('messages').all().order_by('-last_message_at')

    # Sheet 1: Vebsayt
    ws_web = wb.active
    ws_web.title = "Vebsayt Murojaatlari"

    # Sheet 2: Telegram
    ws_tg = wb.create_sheet(title="Telegram Murojaatlari")

    def populate_sheet(ws, queryset, is_telegram=False):
        # Header Row
        ws.append(headers)
        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Data Rows
        row_num = 1
        for s in queryset:
            msg_count = s.messages.count()
            created_str = s.created_at.strftime("%d.%m.%Y %H:%M") if s.created_at else "-"
            last_msg_str = s.last_message_at.strftime("%d.%m.%Y %H:%M") if s.last_message_at else "-"

            name_val = s.visitor_name or s.display_name
            phone_val = s.visitor_phone or "-"
            email_val = s.visitor_email or "-"
            notes_val = s.notes or "-"

            row_data = [
                row_num,
                name_val,
                phone_val,
                email_val,
                created_str,
                last_msg_str,
                msg_count,
                notes_val
            ]
            ws.append(row_data)
            current_row = row_num + 1
            ws.row_dimensions[current_row].height = 20

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if col_idx in [1, 5, 6, 7]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

            row_num += 1

        # Adjust column widths
        col_widths = [6, 26, 20, 26, 20, 20, 15, 35]
        for idx, width in enumerate(col_widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = width

    # Populate both sheets
    web_sessions = all_sessions.exclude(session_id__startswith='tg_')
    tg_sessions = all_sessions.filter(session_id__startswith='tg_')

    populate_sheet(ws_web, web_sessions, is_telegram=False)
    populate_sheet(ws_tg, tg_sessions, is_telegram=True)

    # Save to memory stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Lumos_Murojaatlar_{business.name}_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


